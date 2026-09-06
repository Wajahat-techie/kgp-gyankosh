"""
Document Loader for KGP Gyankosh
Format-aware unified document loader supporting PDF, DOCX, TXT, and scanned image formats.
Automatically detects scanned PDFs and invokes OCR fallback where needed.
Enriches every loaded document with standard metadata: source, page, file_type, and hash.
"""

import os
import hashlib
import logging
from typing import List, Optional
from langchain_core.documents import Document

# Supported document formats
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".png", ".jpg", ".jpeg", ".tiff", ".bmp"}

logger = logging.getLogger("kgp_gyankosh.loader")


def compute_file_hash(file_path: str) -> str:
    """Computes MD5 hash of a file to track modifications for incremental indexing."""
    hasher = hashlib.md5()
    with open(file_path, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def load_txt(file_path: str) -> List[Document]:
    """Loads a plain text document with fallback encoding detection."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    text = ""
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                text = f.read()
            break
        except UnicodeDecodeError:
            continue
            
    if not text.strip():
        logger.warning(f"File {file_path} is empty or unreadable.")
        return []

    file_hash = compute_file_hash(file_path)
    metadata = {
        "source": os.path.basename(file_path),
        "file_path": file_path,
        "file_type": "txt",
        "page": 1,
        "file_hash": file_hash,
    }
    return [Document(page_content=text.strip(), metadata=metadata)]


def load_docx(file_path: str) -> List[Document]:
    """Loads a Microsoft Word document (.docx) including tables."""
    try:
        import docx
        doc = docx.Document(file_path)
        content_parts = []
        
        # Paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                content_parts.append(para.text.strip())
                
        # Tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    content_parts.append(row_text)
                    
        full_text = "\n\n".join(content_parts)
        if not full_text.strip():
            logger.warning(f"DOCX document {file_path} contains no readable text.")
            return []
            
        file_hash = compute_file_hash(file_path)
        metadata = {
            "source": os.path.basename(file_path),
            "file_path": file_path,
            "file_type": "docx",
            "page": 1,
            "file_hash": file_hash,
        }
        return [Document(page_content=full_text, metadata=metadata)]
    except Exception as exc:
        logger.error(f"Error loading DOCX file {file_path}: {exc}", exc_info=True)
        return []


def load_image_document(file_path: str) -> List[Document]:
    """Extracts text from scanned image notices using OCR."""
    from src.ingestion.ocr_loader import extract_text_from_image
    
    text = extract_text_from_image(file_path)
    if not text:
        logger.warning(f"No text extracted via OCR from image: {file_path}")
        return []
        
    file_hash = compute_file_hash(file_path)
    metadata = {
        "source": os.path.basename(file_path),
        "file_path": file_path,
        "file_type": os.path.splitext(file_path)[1].lower().replace(".", ""),
        "page": 1,
        "file_hash": file_hash,
        "is_ocr": True
    }
    return [Document(page_content=text, metadata=metadata)]


def load_pdf(file_path: str) -> List[Document]:
    """
    Loads a PDF document page-by-page.
    If a page has minimal or no text layer (e.g. scanned administrative notice),
    it dynamically falls back to OCR via pdf2image and pytesseract.
    """
    import pypdf
    from src.ingestion.ocr_loader import extract_text_from_scanned_pdf
    
    documents: List[Document] = []
    file_hash = compute_file_hash(file_path)
    base_name = os.path.basename(file_path)
    
    try:
        reader = pypdf.PdfReader(file_path)
        total_pages = len(reader.pages)
        scanned_page_indices = []

        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            # If page text is very sparse (< 30 characters), it's likely a scanned notice/stamp
            if len(text) < 30:
                scanned_page_indices.append(idx)
            else:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": base_name,
                            "file_path": file_path,
                            "file_type": "pdf",
                            "page": idx,
                            "total_pages": total_pages,
                            "file_hash": file_hash,
                            "is_ocr": False
                        }
                    )
                )

        # If scanned pages were detected, attempt OCR fallback
        if scanned_page_indices:
            enable_ocr = os.getenv("ENABLE_OCR", "true").lower() == "true"
            if enable_ocr:
                logger.info(
                    f"PDF '{base_name}' contains {len(scanned_page_indices)} potential scanned page(s). Invoking OCR..."
                )
                ocr_pages = extract_text_from_scanned_pdf(file_path)
                for ocr_res in ocr_pages:
                    p_num = ocr_res.get("page", 1)
                    p_text = ocr_res.get("text", "").strip()
                    if p_num in scanned_page_indices and p_text:
                        documents.append(
                            Document(
                                page_content=p_text,
                                metadata={
                                    "source": base_name,
                                    "file_path": file_path,
                                    "file_type": "pdf",
                                    "page": p_num,
                                    "total_pages": total_pages,
                                    "file_hash": file_hash,
                                    "is_ocr": True
                                }
                            )
                        )
                        logger.info(f"Successfully OCR'd page {p_num} for {base_name}")
            else:
                logger.info("OCR is disabled in settings. Skipping scanned page processing.")

        # Sort documents by page number
        documents.sort(key=lambda d: d.metadata.get("page", 1))
        return documents

    except Exception as exc:
        logger.error(f"Error loading PDF {file_path}: {exc}", exc_info=True)
        return []


def load_single_document(file_path: str) -> List[Document]:
    """Dispatches document loading to the appropriate handler based on file extension."""
    if not os.path.exists(file_path):
        logger.error(f"File does not exist: {file_path}")
        return []

    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".pdf":
        return load_pdf(file_path)
    elif ext == ".docx":
        return load_docx(file_path)
    elif ext == ".txt":
        return load_txt(file_path)
    elif ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return load_image_document(file_path)
    else:
        logger.warning(f"Unsupported file format '{ext}' for file: {file_path}")
        return []


def load_documents_from_directory(directory_path: str) -> List[Document]:
    """
    Recursively scans and loads all supported documents from a directory.
    
    Args:
        directory_path: Directory containing administrative notices.
        
    Returns:
        List of all extracted Document objects with rich metadata.
    """
    if not os.path.exists(directory_path):
        logger.warning(f"Directory does not exist: {directory_path}")
        return []

    all_docs: List[Document] = []
    file_count = 0

    for root, _, files in os.walk(directory_path):
        for file in sorted(files):
            ext = os.path.splitext(file)[1].lower()
            if ext in SUPPORTED_EXTENSIONS:
                full_path = os.path.join(root, file)
                try:
                    docs = load_single_document(full_path)
                    all_docs.extend(docs)
                    file_count += 1
                    logger.info(f"Loaded '{file}': {len(docs)} page(s)/section(s)")
                except Exception as file_err:
                    logger.error(f"Failed to load file '{full_path}': {file_err}")

    logger.info(f"Ingestion complete: {file_count} files processed, {len(all_docs)} total document segments loaded.")
    return all_docs
